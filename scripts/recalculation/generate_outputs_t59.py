#!/usr/bin/env python3
"""Tournament 59 outputs: PointsHistory + WeeklyRanking week 9 + exports.

Pipeline:
1. Read draws 233-236 data from DB (DrawPlayers, Matches)
2. Calculate PointsHistory for tournament 59
3. Insert PointsHistory into DB
4. Calculate WeeklyRanking for week 9 (using full 52-week window)
5. Insert WeeklyRanking into DB
6. Export SQL  -> data/sql/generated/
7. Export Excel -> data/extracts/generated/

Does NOT touch existing files for tournaments 1-58.
"""

from __future__ import annotations

import os
import sys
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Any, Dict, List, Tuple

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from openpyxl import Workbook

from src.modules.db_connection import DatabaseConnection
from src.modules.calculate_points_history import calculate_points_history
from src.modules.calculate_weekly_ranking import calculate_weekly_ranking

# ============================================================
# CONFIG
# ============================================================

TOURNAMENT_ID = 59
CATEGORY_ID = 3  # MT400
DRAW_IDS = (233, 234, 235, 236)
RANKING_YEAR = 2026
RANKING_WEEK = 9

REPO_ROOT = Path(ROOT)
SQL_OUT_DIR = REPO_ROOT / "data" / "sql" / "generated"
XLSX_OUT_DIR = REPO_ROOT / "data" / "extracts" / "generated"


# ============================================================
# UTILITIES
# ============================================================

def _strip_tz(v: Any) -> Any:
    if isinstance(v, datetime) and v.tzinfo is not None:
        return v.replace(tzinfo=None)
    return v


def _to_short_date(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.date()
    return v


def _sql_literal(v: Any) -> str:
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return str(v)
    if isinstance(v, datetime):
        return f"'{v.strftime('%Y-%m-%d %H:%M:%S')}'"
    if isinstance(v, date):
        return f"'{v.strftime('%Y-%m-%d')}'"
    s = str(v).replace("'", "''")
    return f"'{s}'"


def write_sql_inserts(out_path: Path, table_name: str,
                      rows: List[Dict[str, Any]], header_comment: str) -> None:
    if not rows:
        out_path.write_text(
            f"-- {header_comment}\n-- Table: {table_name}\n-- Rows: 0\n",
            encoding="utf-8",
        )
        return
    cols = list(rows[0].keys())
    col_list = ", ".join(cols)
    lines = [
        f"-- {header_comment}",
        f"-- Table: {table_name}",
        f"-- Rows: {len(rows)}",
        "",
    ]
    for r in rows:
        values = ", ".join(_sql_literal(r.get(c)) for c in cols)
        lines.append(f"INSERT INTO {table_name} ({col_list}) VALUES ({values});")
    lines.append("")
    out_path.write_text("\n".join(lines), encoding="utf-8")


def write_xlsx(out_path: Path, sheet_name: str,
               rows: List[Dict[str, Any]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    if not rows:
        wb.save(out_path)
        return
    cols = list(rows[0].keys())
    ws.append(cols)
    for r in rows:
        ws.append([_strip_tz(r.get(c)) for c in cols])
    wb.save(out_path)


# ============================================================
# MAIN
# ============================================================

def main() -> None:
    SQL_OUT_DIR.mkdir(parents=True, exist_ok=True)
    XLSX_OUT_DIR.mkdir(parents=True, exist_ok=True)

    db = DatabaseConnection()
    if not db.connect():
        raise RuntimeError("Database connection failed")

    # ----------------------------------------------------------
    # 1. Fetch tournament 59 metadata
    # ----------------------------------------------------------
    t_row = db.query_params(
        "SELECT * FROM Tournaments WHERE tournament_id = %s", (TOURNAMENT_ID,)
    )
    if not t_row:
        raise RuntimeError(f"Tournament {TOURNAMENT_ID} not found")
    t_info = t_row[0]
    t_end_date = t_info["end_date"]
    if isinstance(t_end_date, datetime):
        t_end_date = t_end_date.date()

    # ----------------------------------------------------------
    # 2. Fetch draws for tournament 59
    # ----------------------------------------------------------
    draws = db.query_params(
        "SELECT * FROM Draws WHERE tournament_id = %s ORDER BY draw_id",
        (TOURNAMENT_ID,),
    )
    draw_id_list = [int(d["draw_id"]) for d in draws]

    # ----------------------------------------------------------
    # 3. Fetch DrawPlayers, DrawSeed, Entries, Matches for t59
    # ----------------------------------------------------------
    placeholders = ",".join(["%s"] * len(draw_id_list))

    draw_players = db.query_params(
        f"SELECT * FROM DrawPlayers WHERE draw_id IN ({placeholders}) ORDER BY draw_id, draw_position",
        tuple(draw_id_list),
    )
    draw_seeds = db.query_params(
        f"SELECT * FROM DrawSeed WHERE draw_id IN ({placeholders}) ORDER BY draw_id, seed_number",
        tuple(draw_id_list),
    )
    entries = db.query_params(
        "SELECT * FROM Entries WHERE tournament_id = %s ORDER BY entry_id",
        (TOURNAMENT_ID,),
    )
    matches = db.query_params(
        f"SELECT * FROM Matches WHERE draw_id IN ({placeholders}) ORDER BY match_id",
        tuple(draw_id_list),
    )

    # Convert match_date to short date
    for m in matches:
        if "match_date" in m:
            m["match_date"] = _to_short_date(m["match_date"])

    # ----------------------------------------------------------
    # 4. Fetch points_rules
    # ----------------------------------------------------------
    pr_rows = db.query("SELECT category_id, stage_result_id, points FROM PointsRules")
    points_rules: Dict[Tuple[int, int], int] = {
        (int(r["category_id"]), int(r["stage_result_id"])): int(r["points"])
        for r in pr_rows
    }

    # ----------------------------------------------------------
    # 5. Calculate PointsHistory for tournament 59
    # ----------------------------------------------------------
    next_ph_row = db.query("SELECT COALESCE(MAX(id), 0) + 1 AS next_id FROM PointsHistory")
    next_ph_id = int(next_ph_row[0]["next_id"])

    # Check for existing t59 PointsHistory (avoid duplicates)
    existing_ph = db.query_params(
        "SELECT COUNT(*) AS cnt FROM PointsHistory WHERE tournament_id = %s",
        (TOURNAMENT_ID,),
    )
    if int(existing_ph[0]["cnt"]) > 0:
        print(f"PointsHistory for tournament {TOURNAMENT_ID} already exists — skipping insert.")
        ph_rows_t59 = [
            dict(r) for r in db.query_params(
                "SELECT * FROM PointsHistory WHERE tournament_id = %s ORDER BY id",
                (TOURNAMENT_ID,),
            )
        ]
    else:
        ph_rows_t59: List[Dict[str, Any]] = []

        for draw in draws:
            did = int(draw["draw_id"])
            ac_id = int(draw["age_category_id"])

            draw_matches = [m for m in matches if int(m["draw_id"]) == did]
            draw_dps = [dp for dp in draw_players if int(dp["draw_id"]) == did]
            if not draw_matches or not draw_dps:
                continue

            player_ids = [int(dp["player_id"]) for dp in draw_dps]

            ph = calculate_points_history(
                draw_id=did,
                tournament_id=TOURNAMENT_ID,
                age_category_id=ac_id,
                category_id=CATEGORY_ID,
                points_rules=points_rules,
                tournament_end_date=t_end_date,
                matches=draw_matches,
                player_ids=player_ids,
                next_ph_id=next_ph_id,
            )

            # Ensure date fields
            for row in ph:
                row["tournament_end_date"] = _to_short_date(row["tournament_end_date"])

            ph_rows_t59.extend(ph)
            if ph:
                next_ph_id = max(int(p["id"]) for p in ph) + 1

        # Insert into DB
        if ph_rows_t59:
            cur = db.connection.cursor()
            for r in ph_rows_t59:
                cur.execute(
                    """
                    INSERT INTO PointsHistory
                      (id, player_id, tournament_id, age_category_id,
                       stage_result_id, points_earned, tournament_end_date, created_at)
                    VALUES (%s, %s, %s, %s, %s, %s, %s, %s)
                    """,
                    (
                        r["id"], r["player_id"], r["tournament_id"],
                        r["age_category_id"], r["stage_result_id"],
                        r["points_earned"], r["tournament_end_date"],
                        r["created_at"],
                    ),
                )
            db.connection.commit()
            cur.close()
            print(f"Inserted {len(ph_rows_t59)} PointsHistory rows for tournament {TOURNAMENT_ID}")

    # ----------------------------------------------------------
    # 6. Calculate WeeklyRanking for week 9
    # ----------------------------------------------------------

    # Fetch ALL PointsHistory (52-week window needs historical data)
    all_ph = [dict(r) for r in db.query("SELECT * FROM PointsHistory ORDER BY id")]

    # Build tournament week metadata
    all_tournaments = db.query("SELECT tournament_id, tournament_year, tournament_week FROM Tournaments")
    tournament_week_meta: Dict[int, Dict[str, int]] = {
        int(t["tournament_id"]): {
            "year": int(t["tournament_year"]),
            "week": int(t["tournament_week"]),
        }
        for t in all_tournaments
    }

    # Player gender map
    all_players = db.query("SELECT player_id, gender_id FROM Players")
    player_gender: Dict[int, int] = {
        int(p["player_id"]): int(p["gender_id"]) for p in all_players
    }

    # Check for existing week 9 ranking
    existing_wr = db.query_params(
        "SELECT COUNT(*) AS cnt FROM WeeklyRanking WHERE ranking_year = %s AND ranking_week = %s",
        (RANKING_YEAR, RANKING_WEEK),
    )

    if int(existing_wr[0]["cnt"]) > 0:
        print(f"WeeklyRanking for {RANKING_YEAR}-W{RANKING_WEEK} already exists — skipping insert.")
        wr_rows = [
            dict(r) for r in db.query_params(
                "SELECT * FROM WeeklyRanking WHERE ranking_year = %s AND ranking_week = %s ORDER BY age_category_id, gender_id, rank_position",
                (RANKING_YEAR, RANKING_WEEK),
            )
        ]
    else:
        wr_rows = calculate_weekly_ranking(
            ranking_year=RANKING_YEAR,
            ranking_week=RANKING_WEEK,
            points_history=all_ph,
            tournaments=tournament_week_meta,
            player_gender=player_gender,
        )

        if wr_rows:
            cur = db.connection.cursor()
            for r in wr_rows:
                cur.execute(
                    """
                    INSERT INTO WeeklyRanking
                      (player_id, age_category_id, gender_id,
                       ranking_year, ranking_week, total_points, rank_position)
                    VALUES (%s, %s, %s, %s, %s, %s, %s)
                    """,
                    (
                        r["player_id"], r["age_category_id"], r["gender_id"],
                        r["ranking_year"], r["ranking_week"],
                        r["total_points"], r["rank_position"],
                    ),
                )
            db.connection.commit()
            cur.close()
            print(f"Inserted {len(wr_rows)} WeeklyRanking rows for {RANKING_YEAR}-W{RANKING_WEEK}")

    # ----------------------------------------------------------
    # 7. Fetch PlayerSuspensions for t59 (if any)
    # ----------------------------------------------------------
    suspensions_t59 = db.query_params(
        "SELECT * FROM PlayerSuspensions WHERE tournament_id = %s ORDER BY suspension_id",
        (TOURNAMENT_ID,),
    )

    db.disconnect()

    # ----------------------------------------------------------
    # 8. Export SQL
    # ----------------------------------------------------------
    header = "Generated by scripts/recalculation/generate_outputs_t59.py"

    write_sql_inserts(
        SQL_OUT_DIR / "entries_t59.sql", "Entries", entries, header,
    )
    write_sql_inserts(
        SQL_OUT_DIR / "drawplayers_t59.sql", "DrawPlayers", draw_players, header,
    )
    write_sql_inserts(
        SQL_OUT_DIR / "drawseed_t59.sql", "DrawSeed", draw_seeds, header,
    )
    write_sql_inserts(
        SQL_OUT_DIR / "matches_t59.sql", "Matches", matches, header,
    )
    write_sql_inserts(
        SQL_OUT_DIR / "playersuspensions_t59.sql", "PlayerSuspensions",
        suspensions_t59, header,
    )
    write_sql_inserts(
        SQL_OUT_DIR / "pointshistory_t59.sql", "PointsHistory", ph_rows_t59, header,
    )
    write_sql_inserts(
        SQL_OUT_DIR / "weeklyranking_year2026_week9.sql", "WeeklyRanking",
        wr_rows, header,
    )

    # _run_all_generated_59.sql
    run_all_lines = [
        "-- Generated data for tournament 59",
        "-- Run AFTER _run_all_generated_1to58.sql",
        "-- Order follows the generation pipeline",
        "",
        "\\i data/sql/generated/entries_t59.sql",
        "\\i data/sql/generated/drawplayers_t59.sql",
        "\\i data/sql/generated/drawseed_t59.sql",
        "\\i data/sql/generated/matches_t59.sql",
        "\\i data/sql/generated/playersuspensions_t59.sql",
        "\\i data/sql/generated/pointshistory_t59.sql",
        "\\i data/sql/generated/weeklyranking_year2026_week9.sql",
    ]
    (SQL_OUT_DIR / "_run_all_generated_59.sql").write_text(
        "\n".join(run_all_lines) + "\n", encoding="utf-8",
    )

    # ----------------------------------------------------------
    # 9. Export Excel
    # ----------------------------------------------------------
    write_xlsx(XLSX_OUT_DIR / "Entries_t59.xlsx", "Entries", entries)
    write_xlsx(XLSX_OUT_DIR / "DrawPlayers_t59.xlsx", "DrawPlayers", draw_players)
    write_xlsx(XLSX_OUT_DIR / "DrawSeed_t59.xlsx", "DrawSeed", draw_seeds)
    write_xlsx(XLSX_OUT_DIR / "Matches_t59.xlsx", "Matches", matches)
    write_xlsx(XLSX_OUT_DIR / "PlayerSuspensions_t59.xlsx", "PlayerSuspensions", suspensions_t59)
    write_xlsx(XLSX_OUT_DIR / "PointsHistory_t59.xlsx", "PointsHistory", ph_rows_t59)
    write_xlsx(XLSX_OUT_DIR / "weeklyranking_year2026_week9.xlsx", "WeeklyRanking", wr_rows)

    # ----------------------------------------------------------
    # Summary
    # ----------------------------------------------------------
    print("\n=== Export complete ===")
    print(f"  Entries:            {len(entries)} rows")
    print(f"  DrawPlayers:        {len(draw_players)} rows")
    print(f"  DrawSeed:           {len(draw_seeds)} rows")
    print(f"  Matches:            {len(matches)} rows")
    print(f"  PlayerSuspensions:  {len(suspensions_t59)} rows")
    print(f"  PointsHistory:      {len(ph_rows_t59)} rows")
    print(f"  WeeklyRanking W9:   {len(wr_rows)} rows")
    print(f"\nSQL  -> {SQL_OUT_DIR}")
    print(f"Excel -> {XLSX_OUT_DIR}")


if __name__ == "__main__":
    main()
